#!/usr/bin/env python3
"""Поиск ниш на YouTube через официальный YouTube Data API v3.

Ищет видео по списку запросов из config.json, находит ролики небольших
каналов с большим числом просмотров и сохраняет отчёт niches_ДАТА.xlsx.

Запуск:  python niche_finder.py            (берёт config.json рядом с программой)
         python niche_finder.py my.json    (другой файл настроек)
"""

from __future__ import annotations

import hashlib
import json
import math
import os
import re
import shutil
import statistics
import sys
import time
import unicodedata
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

try:
    import pandas as pd
    import requests
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    print(f"Не установлена библиотека: {exc.name}")
    print("Установите зависимости командой:  pip install -r requirements.txt")
    sys.exit(1)

BASE_DIR = Path(__file__).resolve().parent
CACHE_DIR = BASE_DIR / "cache"
API_URL = "https://www.googleapis.com/youtube/v3/"
SUGGEST_URL = "https://suggestqueries.google.com/complete/search"
TRANSLATE_CLOUD_URL = "https://translation.googleapis.com/language/translate/v2"
TRANSLATE_FREE_URL = "https://translate.googleapis.com/translate_a/single"

# Стоимость одного вызова в единицах квоты.
COST = {"search": 100, "videos": 1, "channels": 1}
PAGE_SIZE = 50  # максимум maxResults для search.list, а также id в videos.list и channels.list

KEY_PLACEHOLDER = "ВСТАВЬТЕ_СЮДА_КЛЮЧ_API"

DEFAULT_CONFIG = {
    "api_key": KEY_PLACEHOLDER,
    "queries": [
        {"q": "why did the empire collapse", "lang": "en", "region": "US"},
        {"q": "jak ludzie wynaleźli", "lang": "pl", "region": "PL"},
    ],
    "order": "viewCount",
    "published_after_days": 365,
    "video_duration": "any",
    "max_results_per_query": 100,
    "subs_min": 0,
    "subs_max": 20000,
    "views_min": 50000,
    "views_max": 5000000,
    "min_duration_minutes": 4,
    "exclude_shorts": True,
    "keep_hidden_subs": True,
    "fresh_days": 90,
    "fresh_bonus": 1.5,
    "suggest": False,
    "suggest_as_queries": False,
    "suggest_max_per_query": 5,
    "translate_check": False,
    "translate_top_videos": 5,
    "translate_source_langs": ["en"],
    "translate_languages": [
        {"lang": "cs", "region": "CZ"},
        {"lang": "ro", "region": "RO"},
        {"lang": "fr", "region": "FR"},
        {"lang": "hu", "region": "HU"},
        {"lang": "sr", "region": "RS"},
        {"lang": "de", "region": "DE"},
        {"lang": "ru", "region": "RU"},
    ],
    "translate_provider": "auto",
    "serbian_script": "both",
    "translate_min_views": 10000,
    "translate_similarity": 0.5,
    "translate_busy_count": 3,
    "daily_quota": 10000,
    "use_cache": True,
    "output_dir": ".",
}

LANG_NAMES = {
    "cs": "Чешский", "ro": "Румынский", "fr": "Французский", "hu": "Венгерский", "sr": "Сербский",
    "de": "Немецкий", "ru": "Русский", "en": "Английский", "pl": "Польский", "uk": "Украинский",
    "es": "Испанский", "it": "Итальянский", "pt": "Португальский", "sk": "Словацкий", "bg": "Болгарский",
    "hr": "Хорватский", "nl": "Нидерландский", "tr": "Турецкий",
}


class FinderError(Exception):
    """Ошибка, которую нужно показать пользователю и завершить работу."""


class QuotaError(FinderError):
    """Квоты не хватает или она закончилась."""


class QueryError(Exception):
    """Ошибка одного запроса (например, неверный regionCode). Запрос пропускается."""


# ---------------------------------------------------------------- настройки


def load_config(path: Path) -> dict:
    if not path.exists():
        example = BASE_DIR / "config.example.json"
        if example.exists():
            shutil.copyfile(example, path)
        else:
            path.write_text(json.dumps(DEFAULT_CONFIG, ensure_ascii=False, indent=2), encoding="utf-8")
        raise FinderError(
            f"Файл настроек не найден, создан новый: {path}\n"
            "Откройте его, вставьте ключ API в поле api_key и запустите программу снова."
        )
    try:
        user_cfg = json.loads(path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as exc:
        raise FinderError(
            f"Ошибка в {path.name}: строка {exc.lineno}, позиция {exc.colno}: {exc.msg}.\n"
            "Проверьте кавычки, запятые и скобки (в JSON нельзя ставить запятую после последнего элемента)."
        ) from None
    if not isinstance(user_cfg, dict):
        raise FinderError(f"{path.name} должен содержать объект JSON в фигурных скобках {{ }}.")

    cfg = {**DEFAULT_CONFIG, **user_cfg}

    key = str(cfg.get("api_key") or "").strip()
    if not key or key == KEY_PLACEHOLDER:
        key = os.environ.get("YOUTUBE_API_KEY", "").strip()
    if not key:
        raise FinderError(f"Не указан ключ API. Вставьте его в поле api_key файла {path.name}.")
    cfg["api_key"] = key

    if cfg["order"] not in ("viewCount", "relevance"):
        raise FinderError('order должен быть "viewCount" или "relevance".')
    if cfg["video_duration"] not in ("any", "medium", "long"):
        raise FinderError('video_duration должен быть "any", "medium" или "long".')

    if cfg["translate_provider"] not in ("auto", "google_cloud", "google_free"):
        raise FinderError('translate_provider должен быть "auto", "google_cloud" или "google_free".')
    if cfg["serbian_script"] not in ("cyrillic", "latin", "both"):
        raise FinderError('serbian_script должен быть "cyrillic", "latin" или "both".')

    for name in ("published_after_days", "max_results_per_query", "suggest_max_per_query", "daily_quota",
                 "translate_top_videos", "translate_min_views", "translate_similarity", "translate_busy_count"):
        value = cfg[name]
        if value is not None and (not isinstance(value, (int, float)) or isinstance(value, bool) or value < 0):
            raise FinderError(f"{name} должен быть неотрицательным числом (сейчас: {value!r}).")
    for name in ("subs_min", "subs_max", "views_min", "views_max", "min_duration_minutes",
                 "fresh_days", "fresh_bonus"):
        value = cfg[name]
        if value is not None and (not isinstance(value, (int, float)) or isinstance(value, bool)):
            raise FinderError(f"{name} должен быть числом или null (сейчас: {value!r}).")

    cfg["max_results_per_query"] = int(cfg["max_results_per_query"] or PAGE_SIZE)
    if cfg["max_results_per_query"] < 1:
        raise FinderError("max_results_per_query должен быть не меньше 1.")

    queries = []
    raw_queries = cfg.get("queries")
    if not isinstance(raw_queries, list) or not raw_queries:
        raise FinderError("queries должен быть непустым списком запросов.")
    for i, item in enumerate(raw_queries, 1):
        if isinstance(item, str):
            item = {"q": item}
        if not isinstance(item, dict) or not str(item.get("q") or "").strip():
            raise FinderError(f'Запрос №{i} в queries должен выглядеть так: {{"q": "текст", "lang": "en", "region": "US"}}.')
        queries.append({
            "q": str(item["q"]).strip(),
            "lang": str(item.get("lang") or "").strip(),
            "region": str(item.get("region") or "").strip().upper(),
            "source": "config",
        })
    cfg["queries"] = queries

    languages = []
    for i, item in enumerate(cfg.get("translate_languages") or [], 1):
        if isinstance(item, str):
            item = {"lang": item}
        if not isinstance(item, dict) or not str(item.get("lang") or "").strip():
            raise FinderError(f'Язык №{i} в translate_languages должен выглядеть так: {{"lang": "cs", "region": "CZ"}}.')
        code = str(item["lang"]).strip().lower()
        if code not in {x["lang"] for x in languages}:
            languages.append({"lang": code, "region": str(item.get("region") or "").strip().upper()})
    if cfg["translate_check"] and not languages:
        raise FinderError("translate_check включён, но список translate_languages пуст.")
    cfg["translate_languages"] = languages
    cfg["translate_source_langs"] = [str(x).strip().lower() for x in cfg.get("translate_source_langs") or []]
    cfg["translate_top_videos"] = int(cfg["translate_top_videos"] or 0)
    return cfg


# ---------------------------------------------------------------- кэш и квота


class Cache:
    """Ответы API за текущий день лежат в cache/ГГГГ-ММ-ДД/."""

    def __init__(self, enabled: bool):
        self.enabled = enabled
        self.dir = CACHE_DIR / date.today().isoformat()
        if enabled:
            self.dir.mkdir(parents=True, exist_ok=True)

    def _path(self, name: str, params: dict) -> Path:
        raw = json.dumps(params, sort_keys=True, ensure_ascii=False)
        return self.dir / f"{name}_{hashlib.sha1(raw.encode('utf-8')).hexdigest()[:20]}.json"

    def get(self, name: str, params: dict):
        if not self.enabled:
            return None
        path = self._path(name, params)
        if not path.exists():
            return None
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return None

    def put(self, name: str, params: dict, data) -> None:
        if self.enabled:
            self._path(name, params).write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    def load_items(self, name: str) -> dict:
        if not self.enabled:
            return {}
        path = self.dir / f"{name}.json"
        try:
            return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
        except (OSError, json.JSONDecodeError):
            return {}

    def save_items(self, name: str, items: dict) -> None:
        if self.enabled:
            (self.dir / f"{name}.json").write_text(json.dumps(items, ensure_ascii=False), encoding="utf-8")


class Quota:
    """Локальный счётчик потраченной квоты за день (хранится в cache/quota.json)."""

    def __init__(self, limit: int):
        self.limit = int(limit)
        self.path = CACHE_DIR / "quota.json"
        self.today = date.today().isoformat()
        self.used = 0
        self.run_used = 0
        try:
            data = json.loads(self.path.read_text(encoding="utf-8"))
            if data.get("date") == self.today:
                self.used = int(data.get("used", 0))
        except (OSError, ValueError, AttributeError):
            pass

    @property
    def remaining(self) -> int:
        return max(self.limit - self.used, 0)

    def check(self, cost: int) -> None:
        if self.used + cost > self.limit:
            raise QuotaError(
                f"Квоты не хватает: нужно ещё {cost} ед., осталось {self.remaining} из {self.limit}.\n"
                "Квота обновляется в полночь по тихоокеанскому времени (США). "
                "Уже скачанные данные сохранены в кэше."
            )

    def add(self, cost: int) -> None:
        self.used += cost
        self.run_used += cost
        CACHE_DIR.mkdir(parents=True, exist_ok=True)
        self.path.write_text(json.dumps({"date": self.today, "used": self.used}), encoding="utf-8")

    def summary(self) -> str:
        return (f"Квота: потрачено за этот запуск {self.run_used} ед., "
                f"за сегодня около {self.used} из {self.limit} (по локальному счётчику).")


# ---------------------------------------------------------------- YouTube API


def explain_api_error(resp: requests.Response) -> Exception:
    try:
        err = resp.json().get("error", {})
    except ValueError:
        err = {}
    if not isinstance(err, dict):
        err = {}
    message = err.get("message") or resp.text[:300]
    reasons = {e.get("reason") for e in err.get("errors", []) if isinstance(e, dict)}
    reasons |= {d.get("reason") for d in err.get("details", []) if isinstance(d, dict)}

    if "API_KEY_INVALID" in reasons or "keyInvalid" in reasons or "API key not valid" in message:
        return FinderError(
            "Неверный ключ API. Проверьте поле api_key в config.json: "
            "ключ нужно скопировать целиком, без пробелов и кавычек внутри."
        )
    if reasons & {"quotaExceeded", "dailyLimitExceeded"}:
        return QuotaError(
            "Дневная квота YouTube API закончилась (ответ сервера quotaExceeded).\n"
            "Квота обновляется в полночь по тихоокеанскому времени (США). "
            "Уже скачанные данные сохранены в кэше."
        )
    if reasons & {"accessNotConfigured", "SERVICE_DISABLED"}:
        return FinderError(
            "YouTube Data API v3 не включён в проекте Google Cloud, к которому относится ключ.\n"
            "Откройте Google Cloud Console, раздел APIs & Services > Library, найдите "
            "YouTube Data API v3 и нажмите Enable. Включение может занять несколько минут."
        )
    if reasons & {"API_KEY_SERVICE_BLOCKED", "API_KEY_HTTP_REFERRER_BLOCKED", "API_KEY_IP_ADDRESS_BLOCKED",
                  "API_KEY_ANDROID_APP_BLOCKED", "API_KEY_IOS_APP_BLOCKED", "ipRefererBlocked"}:
        return FinderError(
            "Ключ API запрещает этот запрос из-за ограничений ключа.\n"
            "В Google Cloud Console (Credentials > ваш ключ) проверьте раздел API restrictions: "
            "там должен быть разрешён YouTube Data API v3, а Application restrictions должно быть None."
        )
    if resp.status_code in (400, 404):
        return QueryError(f"YouTube API отклонил запрос ({resp.status_code}): {message}")
    return FinderError(f"Ошибка YouTube API ({resp.status_code}): {message}")


class YouTube:
    def __init__(self, api_key: str, cache: Cache, quota: Quota):
        self.api_key = api_key
        self.cache = cache
        self.quota = quota
        self.session = requests.Session()

    def _request(self, endpoint: str, params: dict) -> dict:
        cost = COST[endpoint]
        self.quota.check(cost)
        last_exc: Exception | None = None
        for attempt in range(3):
            try:
                resp = self.session.get(API_URL + endpoint, params={**params, "key": self.api_key}, timeout=30)
            except requests.ConnectionError:
                raise FinderError(
                    "Нет подключения к интернету или сервер YouTube недоступен. "
                    "Проверьте соединение (и прокси/VPN, если используете) и запустите снова."
                ) from None
            except requests.Timeout as exc:
                last_exc = exc
                time.sleep(2 * (attempt + 1))
                continue
            if resp.status_code in (500, 502, 503, 504):
                last_exc = FinderError(f"Сервер YouTube временно недоступен (код {resp.status_code}).")
                time.sleep(2 * (attempt + 1))
                continue
            if resp.status_code != 200:
                raise explain_api_error(resp)
            self.quota.add(cost)
            return resp.json()
        if isinstance(last_exc, requests.Timeout):
            raise FinderError("Сервер YouTube не отвечает (истекло время ожидания). Попробуйте позже.")
        raise last_exc or FinderError("Не удалось выполнить запрос к YouTube API.")

    @staticmethod
    def search_params(query: dict, cfg: dict, published_after: str | None, page_token: str | None) -> dict:
        params = {
            "part": "snippet",
            "type": "video",
            "q": query["q"],
            "order": cfg["order"],
            "maxResults": PAGE_SIZE,
        }
        if query["lang"]:
            params["relevanceLanguage"] = query["lang"]
        if query["region"]:
            params["regionCode"] = query["region"]
        if published_after:
            params["publishedAfter"] = published_after
        if cfg["video_duration"] != "any":
            params["videoDuration"] = cfg["video_duration"]
        if page_token:
            params["pageToken"] = page_token
        return params

    def search_cost_estimate(self, query: dict, cfg: dict, published_after: str | None) -> int:
        """Сколько единиц уйдёт на запрос с учётом страниц, уже лежащих в кэше."""
        pages = math.ceil(cfg["max_results_per_query"] / PAGE_SIZE)
        token = None
        for page in range(pages):
            data = self.cache.get("search", self.search_params(query, cfg, published_after, token))
            if data is None:
                return (pages - page) * COST["search"]
            token = data.get("nextPageToken")
            if not token or not data.get("items"):
                return 0
        return 0

    def search(self, query: dict, cfg: dict, published_after: str | None) -> tuple[list[str], int, int]:
        """Возвращает (id видео, страниц из сети, страниц из кэша)."""
        limit = cfg["max_results_per_query"]
        ids: list[str] = []
        seen: set[str] = set()
        token = None
        net_pages = cached_pages = 0
        while len(ids) < limit:
            params = self.search_params(query, cfg, published_after, token)
            data = self.cache.get("search", params)
            if data is None:
                data = self._request("search", params)
                self.cache.put("search", params, data)
                net_pages += 1
            else:
                cached_pages += 1
            items = data.get("items") or []
            for item in items:
                vid = (item.get("id") or {}).get("videoId")
                if vid and vid not in seen:
                    seen.add(vid)
                    ids.append(vid)
            token = data.get("nextPageToken")
            if not token or not items:
                break
        return ids[:limit], net_pages, cached_pages

    def _fetch_items(self, endpoint: str, part: str, ids: list[str]) -> dict:
        store = self.cache.load_items(endpoint)
        missing = [i for i in dict.fromkeys(ids) if i not in store]
        for start in range(0, len(missing), PAGE_SIZE):
            batch = missing[start:start + PAGE_SIZE]
            try:
                data = self._request(endpoint, {"part": part, "id": ",".join(batch), "maxResults": PAGE_SIZE})
            except QueryError as exc:
                raise FinderError(str(exc)) from None
            for item in data.get("items") or []:
                store[item["id"]] = item
            for vid in batch:
                store.setdefault(vid, None)  # удалённые или приватные: больше не запрашиваем
            self.cache.save_items(endpoint, store)
        return {i: store.get(i) for i in ids}

    def videos(self, ids: list[str]) -> dict:
        return self._fetch_items("videos", "snippet,statistics,contentDetails", ids)

    def channels(self, ids: list[str]) -> dict:
        return self._fetch_items("channels", "snippet,statistics", ids)


# ---------------------------------------------------------------- подсказки


def fetch_suggestions(session: requests.Session, cache: Cache, query: dict) -> list[str]:
    params = {"client": "firefox", "ds": "yt", "hl": query["lang"] or "en", "q": query["q"]}
    cached = cache.get("suggest", params)
    if cached is not None:
        return cached
    resp = session.get(SUGGEST_URL, params=params, timeout=15, headers={"User-Agent": "Mozilla/5.0"})
    resp.raise_for_status()
    match = re.search(r"charset=([\w.-]+)", resp.headers.get("Content-Type", ""), re.I)
    encodings = [match.group(1)] if match else []
    encodings += ["utf-8", "cp1252"]
    text = None
    for enc in encodings:
        try:
            text = resp.content.decode(enc)
            break
        except (UnicodeDecodeError, LookupError):
            continue
    if text is None:
        text = resp.content.decode("utf-8", errors="replace")
    data = json.loads(text)
    suggestions = [s for s in data[1] if isinstance(s, str)] if isinstance(data, list) and len(data) > 1 else []
    cache.put("suggest", params, suggestions)
    return suggestions


# ---------------------------------------------------------------- перевод

CLOUD_TRANSLATE_HELP = (
    "Чтобы переводить через официальный Cloud Translation API, в Google Cloud Console того же проекта:\n"
    "  1) APIs & Services > Library: найдите Cloud Translation API и нажмите Enable;\n"
    "  2) Billing: подключите платёжный аккаунт (API тарифицируется по числу символов, "
    "есть бесплатный месячный объём, условия на странице цен Cloud Translation);\n"
    "  3) если у ключа включены API restrictions, добавьте в список Cloud Translation API."
)

SR_CYR_TO_LAT = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "ђ": "đ", "е": "e", "ж": "ž", "з": "z", "и": "i",
    "ј": "j", "к": "k", "л": "l", "љ": "lj", "м": "m", "н": "n", "њ": "nj", "о": "o", "п": "p", "р": "r",
    "с": "s", "т": "t", "ћ": "ć", "у": "u", "ф": "f", "х": "h", "ц": "c", "ч": "č", "џ": "dž", "ш": "š",
}


def sr_to_latin(text: str) -> str:
    out = []
    for ch in text:
        lat = SR_CYR_TO_LAT.get(ch.lower())
        if lat is None:
            out.append(ch)
        else:
            out.append(lat[0].upper() + lat[1:] if ch.isupper() else lat)
    return "".join(out)


def clean_title(title: str) -> str:
    """Убирает хэштеги и эмодзи, чтобы переводилась только суть названия."""
    text = re.sub(r"#\w+", " ", title)
    text = "".join(ch for ch in text if unicodedata.category(ch) not in ("So", "Sk", "Cs", "Co"))
    text = re.sub(r"\s+", " ", text).strip(" |-:")
    return text or title


def title_stems(text: str, lang: str) -> set[str]:
    """Слова названия без регистра и диакритики, обрезанные до 5 букв (грубый учёт падежей)."""
    text = text.lower()
    if lang == "sr":
        text = sr_to_latin(text)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return {w[:5] for w in re.findall(r"\w+", text) if len(w) >= 3}


def similarity(query_stems: set[str], title: str, lang: str) -> float:
    """Какая доля слов переведённого названия встречается в названии найденного ролика."""
    if not query_stems:
        return 0.0
    return len(query_stems & title_stems(title, lang)) / len(query_stems)


class TranslateUnavailable(Exception):
    pass


class Translator:
    def __init__(self, provider: str, api_key: str, cache: Cache, session: requests.Session):
        self.provider = provider
        self.api_key = api_key
        self.cache = cache
        self.session = session
        self.store = cache.load_items("translations")

    def translate(self, texts: list[str], target: str) -> list[str]:
        missing = [t for t in dict.fromkeys(texts) if f"{target}|{t}" not in self.store]
        if missing:
            for text, result in zip(missing, self._translate(missing, target)):
                self.store[f"{target}|{text}"] = result
            self.cache.save_items("translations", self.store)
        return [self.store[f"{target}|{t}"] for t in texts]

    def _translate(self, texts: list[str], target: str) -> list[str]:
        if self.provider in ("auto", "google_cloud"):
            try:
                return self._cloud(texts, target)
            except TranslateUnavailable as exc:
                if self.provider == "google_cloud":
                    raise FinderError(f"Cloud Translation API недоступен: {exc}\n{CLOUD_TRANSLATE_HELP}") from None
                print(f"  Cloud Translation API недоступен ({exc}).\n"
                      "  Перевожу через бесплатный переводчик Google (неофициальный, может блокировать частые запросы).")
                self.provider = "google_free"
        try:
            return [self._free(t, target) for t in texts]
        except TranslateUnavailable as exc:
            raise FinderError(f"Не удалось перевести названия: {exc}\n{CLOUD_TRANSLATE_HELP}") from None

    def _cloud(self, texts: list[str], target: str) -> list[str]:
        result = []
        for start in range(0, len(texts), 100):
            batch = texts[start:start + 100]
            try:
                resp = self.session.post(
                    TRANSLATE_CLOUD_URL, params={"key": self.api_key},
                    json={"q": batch, "target": target, "format": "text"}, timeout=30,
                )
            except requests.ConnectionError:
                raise FinderError("Нет подключения к интернету. Проверьте соединение и запустите снова.") from None
            except requests.Timeout:
                raise TranslateUnavailable("сервер не ответил вовремя") from None
            if resp.status_code != 200:
                try:
                    message = resp.json()["error"]["message"]
                except (ValueError, KeyError, TypeError):
                    message = f"код ответа {resp.status_code}"
                raise TranslateUnavailable(message)
            result += [t.get("translatedText", "") for t in resp.json()["data"]["translations"]]
        return result

    def _free(self, text: str, target: str) -> str:
        params = {"client": "gtx", "sl": "auto", "tl": target, "dt": "t", "q": text}
        try:
            resp = self.session.get(TRANSLATE_FREE_URL, params=params, timeout=15,
                                    headers={"User-Agent": "Mozilla/5.0"})
        except requests.ConnectionError:
            raise FinderError("Нет подключения к интернету. Проверьте соединение и запустите снова.") from None
        except requests.Timeout:
            raise TranslateUnavailable("бесплатный переводчик не ответил вовремя") from None
        if resp.status_code != 200:
            raise TranslateUnavailable(f"бесплатный переводчик ответил кодом {resp.status_code}")
        try:
            data = resp.json()
            translated = "".join(seg[0] for seg in data[0] if seg and isinstance(seg[0], str)).strip()
        except (ValueError, TypeError, IndexError):
            raise TranslateUnavailable("бесплатный переводчик вернул неожиданный ответ") from None
        time.sleep(0.3)
        return translated or text


# ---------------------------------------------------------------- расчёты

DURATION_RE = re.compile(
    r"^P(?:(\d+)W)?(?:(\d+)D)?(?:T(?:(\d+)H)?(?:(\d+)M)?(?:(\d+(?:\.\d+)?)S)?)?$"
)


def parse_duration(value: str) -> int:
    """ISO 8601 (PT1H2M3S) -> секунды."""
    match = DURATION_RE.match(value or "")
    if not match:
        return 0
    weeks, days, hours, minutes, seconds = match.groups()
    return int(
        int(weeks or 0) * 604800 + int(days or 0) * 86400 + int(hours or 0) * 3600
        + int(minutes or 0) * 60 + float(seconds or 0)
    )


def parse_time(value: str) -> datetime:
    return datetime.strptime(value[:19], "%Y-%m-%dT%H:%M:%S").replace(tzinfo=timezone.utc)


def in_range(value, low, high) -> bool:
    if low is not None and value < low:
        return False
    if high is not None and value > high:
        return False
    return True


def build_rows(hits: dict, queries: list[dict], videos: dict, channels: dict, cfg: dict) -> list[dict]:
    now = datetime.now(timezone.utc)
    rows = []
    for vid, query_idx in hits.items():
        video = videos.get(vid)
        if not video:
            continue
        sn = video.get("snippet") or {}
        st = video.get("statistics") or {}
        cd = video.get("contentDetails") or {}
        if sn.get("liveBroadcastContent", "none") != "none":
            continue  # трансляции и премьеры
        channel = channels.get(sn.get("channelId"))
        if not channel:
            continue

        ch_stats = channel.get("statistics") or {}
        hidden = bool(ch_stats.get("hiddenSubscriberCount")) or "subscriberCount" not in ch_stats
        subs = None if hidden else int(ch_stats.get("subscriberCount", 0))
        views = int(st.get("viewCount", 0))
        published = parse_time(sn["publishedAt"])
        age_days = max((now - published).total_seconds() / 86400, 1.0)
        seconds = parse_duration(cd.get("duration", ""))

        fresh = age_days <= (cfg["fresh_days"] or 0)
        ratio = None if hidden else views / max(subs, 1)
        score = None
        if ratio is not None:
            bonus = cfg["fresh_bonus"] if fresh else 1.0
            score = math.log10(max(views, 1)) * math.log10(ratio + 1) * bonus

        text = " ".join([sn.get("title", ""), sn.get("description", ""), " ".join(sn.get("tags") or [])]).lower()
        is_short = seconds <= 60 or (seconds <= 180 and "#short" in text)

        rows.append({
            "id": vid,
            "query_idx": query_idx,
            "title": sn.get("title", ""),
            "channel_id": sn.get("channelId", ""),
            "channel": (channel.get("snippet") or {}).get("title") or sn.get("channelTitle", ""),
            "channel_videos": int(ch_stats.get("videoCount", 0)),
            "hidden": hidden,
            "subs": subs,
            "views": views,
            "ratio": ratio,
            "vpd": views / age_days,
            "minutes": seconds / 60,
            "published": published,
            "fresh": fresh,
            "score": score,
            "is_short": is_short,
        })
    return rows


def passes_base(row: dict, cfg: dict) -> bool:
    if not in_range(row["views"], cfg["views_min"], cfg["views_max"]):
        return False
    if cfg["min_duration_minutes"] is not None and row["minutes"] < cfg["min_duration_minutes"]:
        return False
    if cfg["exclude_shorts"] and row["is_short"]:
        return False
    return True


def passes_subs(row: dict, cfg: dict) -> bool:
    if row["hidden"]:
        return bool(cfg["keep_hidden_subs"])
    return in_range(row["subs"], cfg["subs_min"], cfg["subs_max"])


def sort_key(row: dict):
    return (row["score"] is not None, row["score"] or 0, row["views"])


# ---------------------------------------------------------------- проверка на других языках

VERDICT_FREE, VERDICT_FEW, VERDICT_BUSY = "свободно", "мало конкурентов", "занято"


def lang_name(code: str) -> str:
    return LANG_NAMES.get(code, code)


def translation_check(yt: YouTube, translator: Translator, final_rows: list[dict], queries: list[dict],
                      cfg: dict, quota: Quota) -> tuple[list[dict], list[dict]]:
    """Переводит названия лучших видео и ищет похожие ролики на других языках.

    Возвращает (исходные видео, результаты по парам видео + язык).
    """
    src_langs = set(cfg["translate_source_langs"])
    sources = [
        r for r in final_rows
        if r["score"] is not None
        and (not src_langs or any(queries[i]["lang"].lower() in src_langs for i in r["query_idx"]))
    ][:cfg["translate_top_videos"]]
    if not sources:
        print("\nПроверка на других языках: нет подходящих видео (проверьте translate_source_langs).")
        return [], []
    languages = [l for l in cfg["translate_languages"]]
    print(f"\nПроверка на других языках: {len(sources)} видео, языки: "
          + ", ".join(lang_name(l["lang"]) for l in languages))

    titles = [clean_title(r["title"]) for r in sources]
    variants: dict[tuple[int, str], list[str]] = {}
    for lang in languages:
        code = lang["lang"]
        print(f"Перевод названий: {lang_name(code)}")
        for k, text in enumerate(translator.translate(titles, code)):
            options = [text]
            if code == "sr":
                latin = sr_to_latin(text)
                options = {"cyrillic": [text], "latin": [latin],
                           "both": list(dict.fromkeys([text, latin]))}[cfg["serbian_script"]]
            variants[(k, code)] = options

    check_cfg = {"order": "relevance", "video_duration": "any", "max_results_per_query": PAGE_SIZE}

    def as_query(text: str, lang: dict) -> dict:
        return {"q": text, "lang": lang["lang"], "region": lang["region"], "source": "translate"}

    # Сколько видео можно проверить на оставшуюся квоту.
    allowed, planned = 0, 0
    for k in range(len(sources)):
        searches = [as_query(v, l) for l in languages for v in variants[(k, l["lang"])]]
        cost = sum(yt.search_cost_estimate(q, check_cfg, None) for q in searches)
        cost += len(searches) * COST["videos"]
        if planned + cost > quota.remaining:
            break
        planned += cost
        allowed += 1
    if allowed == 0:
        print(f"Квоты не хватает даже на одно видео (осталось {quota.remaining} ед.). Проверка пропущена.")
        return sources, []
    if allowed < len(sources):
        print(f"Квоты хватает только на {allowed} из {len(sources)} видео, остальные пропущены.")
        sources = sources[:allowed]
    print(f"Оценка расхода квоты на проверку: до {planned} ед.")

    found: dict[tuple[int, str], list[str]] = {}
    for k, row in enumerate(sources):
        print(f"Видео {k + 1} из {len(sources)}: {short(titles[k])}")
        for lang in languages:
            code = lang["lang"]
            ids: list[str] = []
            for text in variants[(k, code)]:
                try:
                    got, _, _ = yt.search(as_query(text, lang), check_cfg, None)
                except QueryError as exc:
                    print(f"  {lang_name(code)}: запрос отклонён: {exc}")
                    continue
                ids += [i for i in got if i not in ids]
            found[(k, code)] = ids
        print(f"  квота за запуск: {quota.run_used} ед.")

    all_ids = list(dict.fromkeys(i for ids in found.values() for i in ids))
    print(f"Загрузка данных о {len(all_ids)} найденных видео...")
    videos = yt.videos(all_ids)

    results = []
    threshold = cfg["translate_similarity"] or 0
    min_views = cfg["translate_min_views"] or 0
    busy = max(int(cfg["translate_busy_count"] or 1), 1)
    for (k, code), ids in found.items():
        stems = [title_stems(v, code) for v in variants[(k, code)]]
        similar = []
        for vid in ids:
            video = videos.get(vid)
            if not video or vid == sources[k]["id"]:
                continue
            title = (video.get("snippet") or {}).get("title", "")
            sim = max(similarity(s, title, code) for s in stems)
            if sim >= threshold:
                views = int((video.get("statistics") or {}).get("viewCount", 0))
                similar.append({"id": vid, "title": title, "views": views, "sim": sim})
        competitors = [m for m in similar if m["views"] >= min_views]
        best = max(competitors or similar, key=lambda m: m["views"]) if similar else None
        if not competitors:
            verdict = VERDICT_FREE
        elif len(competitors) < busy:
            verdict = VERDICT_FEW
        else:
            verdict = VERDICT_BUSY
        results.append({
            "k": k, "lang": code,
            "region": next(l["region"] for l in languages if l["lang"] == code),
            "translation": " / ".join(variants[(k, code)]),
            "found": len(ids), "similar": len(similar), "competitors": len(competitors),
            "best": best, "verdict": verdict,
        })
    return sources, results


# ---------------------------------------------------------------- таблицы

VIDEO_COLUMNS = ["Запрос", "Язык", "Название видео", "Ссылка на видео", "Канал", "Ссылка на канал",
                 "Подписчики", "Просмотры", "Ratio", "Просмотров в день", "Длительность, мин",
                 "Дата публикации", "Score"]


def video_url(vid: str) -> str:
    return f"https://www.youtube.com/watch?v={vid}"


def channel_url(cid: str) -> str:
    return f"https://www.youtube.com/channel/{cid}"


def round_or_none(value, digits):
    return None if value is None else round(value, digits)


def make_video_table(rows: list[dict], queries: list[dict]) -> pd.DataFrame:
    records = []
    for r in rows:
        qs = [queries[i] for i in r["query_idx"]]
        records.append([
            "; ".join(q["q"] for q in qs),
            ", ".join(dict.fromkeys(q["lang"] for q in qs if q["lang"])),
            r["title"],
            video_url(r["id"]),
            r["channel"],
            channel_url(r["channel_id"]),
            "скрыто" if r["hidden"] else r["subs"],
            r["views"],
            round_or_none(r["ratio"], 2),
            round(r["vpd"]),
            round(r["minutes"], 1),
            r["published"].date(),
            round_or_none(r["score"], 3),
        ])
    return pd.DataFrame(records, columns=VIDEO_COLUMNS)


def make_niche_table(all_rows: list[dict], final_ids: set[str], queries: list[dict],
                     failed: dict[int, str], cfg: dict) -> pd.DataFrame:
    fresh_days = cfg["fresh_days"] or 0
    records = []
    for i, q in enumerate(queries):
        found = [r for r in all_rows if i in r["query_idx"]]
        base = [r for r in found if r["base_ok"]]
        final = [r for r in base if r["id"] in final_ids]
        views = [r["views"] for r in final]
        ratios = [r["ratio"] for r in final if r["ratio"] is not None]
        scores = sorted((r["score"] for r in final if r["score"] is not None), reverse=True)
        top = scores[:3]
        records.append([
            q["q"], q["lang"], q["region"],
            "подсказка" if q["source"] == "suggest" else "config",
            len(found),
            len(base),
            len(final),
            round(len(final) / len(base) * 100, 1) if base else None,
            round(statistics.mean(views)) if views else None,
            round(statistics.median(views)) if views else None,
            round(max(ratios), 2) if ratios else None,
            sum(1 for r in final if r["fresh"]),
            round(sum(top) / len(top), 3) if top else 0.0,
            failed.get(i, ""),
        ])
    df = pd.DataFrame(records, columns=[
        "Запрос", "Язык", "Регион", "Источник", "Найдено видео",
        "Прошли фильтры (без подписчиков)", "Из них на каналах в диапазоне подписчиков",
        "Доля малых каналов, %", "Средние просмотры", "Медианные просмотры", "Лучший ratio",
        f"Видео за {fresh_days} дн.", "Оценка ниши", "Ошибка",
    ])
    return df.sort_values("Оценка ниши", ascending=False, kind="stable")


def make_channel_table(rows: list[dict]) -> pd.DataFrame:
    best: dict[str, dict] = {}
    counts: dict[str, int] = {}
    for r in rows:  # rows уже отсортированы по score, первый ролик канала и есть лучший
        counts[r["channel_id"]] = counts.get(r["channel_id"], 0) + 1
        best.setdefault(r["channel_id"], r)
    records = [[
        r["channel"], channel_url(cid), "скрыто" if r["hidden"] else r["subs"], r["channel_videos"],
        counts[cid], r["title"], video_url(r["id"]), r["views"], round_or_none(r["ratio"], 2),
        round_or_none(r["score"], 3),
    ] for cid, r in best.items()]
    return pd.DataFrame(records, columns=[
        "Канал", "Ссылка на канал", "Подписчики", "Видео на канале", "Найдено роликов",
        "Лучший ролик", "Ссылка на ролик", "Просмотры лучшего", "Ratio лучшего", "Score лучшего",
    ])


def make_suggest_table(suggestions: list[tuple[dict, str, bool]]) -> pd.DataFrame:
    return pd.DataFrame(
        [[q["q"], q["lang"], q["region"], s, "да" if used else "нет"] for q, s, used in suggestions],
        columns=["Исходный запрос", "Язык", "Регион", "Подсказка", "Прогнана как запрос"],
    )


def make_languages_table(sources: list[dict], results: list[dict], languages: list[dict]) -> pd.DataFrame:
    """Сводка: строка на исходное видео, столбец на язык."""
    by_key = {(r["k"], r["lang"]): r for r in results}
    records = []
    for k, row in enumerate(sources):
        cells = []
        for lang in languages:
            res = by_key.get((k, lang["lang"]))
            if res is None:
                cells.append("не проверено")
            elif res["verdict"] == VERDICT_FREE:
                cells.append(VERDICT_FREE)
            else:
                cells.append(f"{res['verdict']} ({res['competitors']})")
        free = sum(1 for c in cells if c == VERDICT_FREE)
        records.append([row["title"], video_url(row["id"]), row["views"], round_or_none(row["ratio"], 2),
                        round_or_none(row["score"], 3), free, *cells])
    df = pd.DataFrame(records, columns=["Исходное видео", "Ссылка на видео", "Просмотры", "Ratio", "Score",
                                        "Свободных языков", *[lang_name(l["lang"]) for l in languages]])
    return df.sort_values(["Свободных языков", "Score"], ascending=False, kind="stable")


def make_translations_table(sources: list[dict], results: list[dict], languages: list[dict]) -> pd.DataFrame:
    order = {l["lang"]: i for i, l in enumerate(languages)}
    records = []
    for r in sorted(results, key=lambda x: (x["k"], order.get(x["lang"], 0))):
        src = sources[r["k"]]
        best = r["best"]
        records.append([
            src["title"], video_url(src["id"]), lang_name(r["lang"]), r["region"], r["translation"],
            r["found"], r["similar"], r["competitors"],
            best["title"] if best else None, video_url(best["id"]) if best else None,
            best["views"] if best else None, round(best["sim"] * 100) if best else None, r["verdict"],
        ])
    return pd.DataFrame(records, columns=[
        "Исходное видео", "Ссылка на видео", "Язык", "Регион", "Перевод названия", "Найдено в поиске",
        "Похожих названий", "Конкурентов", "Самый популярный похожий", "Ссылка на похожий",
        "Его просмотры", "Сходство, %", "Вердикт",
    ])


HEADER_FILL = PatternFill("solid", fgColor="DDE3EA")
GREEN = PatternFill("solid", fgColor="C6EFCE")
BRIGHT_GREEN = PatternFill("solid", fgColor="5BD75B")
VERDICT_FILLS = {
    VERDICT_FREE: PatternFill("solid", fgColor="C6EFCE"),
    VERDICT_FEW: PatternFill("solid", fgColor="FFEB9C"),
    VERDICT_BUSY: PatternFill("solid", fgColor="FFC7CE"),
}

NUMBER_FORMATS = {
    "Подписчики": "#,##0", "Просмотры": "#,##0", "Ratio": "0.00", "Просмотров в день": "#,##0",
    "Длительность, мин": "0.0", "Дата публикации": "yyyy-mm-dd", "Score": "0.000",
    "Средние просмотры": "#,##0", "Медианные просмотры": "#,##0", "Лучший ratio": "0.00",
    "Оценка ниши": "0.000", "Просмотры лучшего": "#,##0", "Ratio лучшего": "0.00",
    "Score лучшего": "0.000", "Видео на канале": "#,##0", "Его просмотры": "#,##0",
}
LINK_COLUMNS = {"Ссылка на видео", "Ссылка на канал", "Ссылка на ролик", "Ссылка на похожий"}


def format_sheet(ws, df: pd.DataFrame, highlight: str | None = None) -> None:
    columns = list(df.columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions

    for col_idx, name in enumerate(columns, 1):
        letter = get_column_letter(col_idx)
        values = [str(v) for v in df[name].tolist() if v is not None and v == v]
        width = max([len(name) * 0.9] + [len(v) for v in values[:500]]) + 2
        if name in LINK_COLUMNS:
            width = 44
        ws.column_dimensions[letter].width = min(max(width, 8), 60)
        fmt = NUMBER_FORMATS.get(name)
        for row in range(2, len(df) + 2):
            cell = ws.cell(row=row, column=col_idx)
            if fmt and not isinstance(cell.value, str):
                cell.number_format = fmt
            if name in LINK_COLUMNS and cell.value:
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"

    if highlight == "ratio" and "Ratio" in columns:
        for row_idx, ratio in enumerate(df["Ratio"].tolist(), start=2):
            if ratio is None or ratio != ratio:
                continue
            fill = BRIGHT_GREEN if ratio > 50 else GREEN if ratio > 10 else None
            if fill:
                for col_idx in range(1, len(columns) + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = fill

    if highlight == "verdict":
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    fill = VERDICT_FILLS.get(cell.value.split(" (")[0])
                    if fill:
                        cell.fill = fill


def write_excel(path: Path, sheets: list[tuple[str, pd.DataFrame, str | None]]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df, highlight in sheets:
            df.to_excel(writer, sheet_name=name, index=False)
            format_sheet(writer.sheets[name], df, highlight)


def save_report(cfg: dict, sheets: list[tuple[str, pd.DataFrame, str | None]]) -> Path:
    out_dir = Path(cfg["output_dir"] or ".")
    if not out_dir.is_absolute():
        out_dir = BASE_DIR / out_dir
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"niches_{date.today().isoformat()}.xlsx"
    try:
        write_excel(path, sheets)
    except PermissionError:
        # Файл за сегодня открыт в Excel: сохраняем рядом под другим именем.
        path = out_dir / f"niches_{datetime.now().strftime('%Y-%m-%d_%H%M%S')}.xlsx"
        write_excel(path, sheets)
    return path


# ---------------------------------------------------------------- основной сценарий


def short(text: str, size: int = 60) -> str:
    return text if len(text) <= size else text[:size - 3] + "..."


def run(config_path: Path) -> None:
    cfg = load_config(config_path)
    cache = Cache(bool(cfg["use_cache"]))
    quota = Quota(cfg["daily_quota"])
    yt = YouTube(cfg["api_key"], cache, quota)
    queries: list[dict] = cfg["queries"]

    published_after = None
    if cfg["published_after_days"]:
        # Округляем до начала суток, чтобы повторный запуск в тот же день попадал в кэш.
        start = datetime.now(timezone.utc) - timedelta(days=cfg["published_after_days"])
        published_after = start.strftime("%Y-%m-%dT00:00:00Z")

    # 1. Подсказки (не тратят квоту API).
    suggestions: list[tuple[dict, str, bool]] = []
    if cfg["suggest"]:
        known = {q["q"].lower() for q in queries}
        extra = []
        base_queries = list(queries)
        for n, q in enumerate(base_queries, 1):
            print(f"Подсказки {n} из {len(base_queries)}: {short(q['q'])}")
            try:
                found = fetch_suggestions(yt.session, cache, q)
            except requests.ConnectionError:
                raise FinderError("Нет подключения к интернету. Проверьте соединение и запустите снова.") from None
            except (requests.RequestException, ValueError) as exc:
                print(f"  не удалось получить подсказки: {exc}")
                continue
            added = 0
            for s in found:
                use = (cfg["suggest_as_queries"] and added < cfg["suggest_max_per_query"]
                       and s.lower() not in known)
                if use:
                    known.add(s.lower())
                    extra.append({"q": s, "lang": q["lang"], "region": q["region"], "source": "suggest"})
                    added += 1
                suggestions.append((q, s, use))
            print(f"  подсказок: {len(found)}" + (f", добавлено как запросы: {added}" if added else ""))
        queries = queries + extra

    # 2. Проверка квоты до начала работы.
    search_cost = sum(yt.search_cost_estimate(q, cfg, published_after) for q in queries)
    max_ids = len(queries) * cfg["max_results_per_query"]
    lookup_cost = 2 * math.ceil(max_ids / PAGE_SIZE) if search_cost else 0
    estimate = search_cost + lookup_cost
    print(f"\nЗапросов: {len(queries)}. Оценка расхода квоты: до {estimate} ед., "
          f"осталось {quota.remaining} из {quota.limit}.")
    if cfg["translate_check"]:
        per_video = sum(2 if (l["lang"] == "sr" and cfg["serbian_script"] == "both") else 1
                        for l in cfg["translate_languages"])
        extra = cfg["translate_top_videos"] * per_video * (COST["search"] + COST["videos"])
        print(f"Проверка на других языках потратит ещё до {extra} ед. "
              "(если квоты не хватит, проверим меньше видео).")
    if estimate > quota.remaining:
        raise QuotaError(
            f"Квоты не хватает: для запуска нужно до {estimate} ед., а осталось {quota.remaining}.\n"
            f"Каждая страница поиска (до 50 видео) стоит {COST['search']} ед. "
            "Уменьшите число запросов или max_results_per_query в config.json, либо запустите "
            "программу после сброса квоты (полночь по тихоокеанскому времени США)."
        )

    # 3. Поиск.
    hits: dict[str, list[int]] = {}
    failed: dict[int, str] = {}
    for i, q in enumerate(queries):
        where = "/".join(x for x in (q["lang"], q["region"]) if x)
        print(f"Запрос {i + 1} из {len(queries)}: {short(q['q'])}" + (f" [{where}]" if where else ""))
        try:
            ids, net_pages, cached_pages = yt.search(q, cfg, published_after)
        except QueryError as exc:
            failed[i] = str(exc)
            print(f"  пропущен: {exc}")
            continue
        for vid in ids:
            hits.setdefault(vid, [])
            if i not in hits[vid]:
                hits[vid].append(i)
        source = f"страниц из сети: {net_pages}, из кэша: {cached_pages}"
        print(f"  найдено видео: {len(ids)} ({source}); квота за запуск: {quota.run_used} ед.")

    if not hits:
        print("\nНичего не найдено. Попробуйте другие запросы или ослабьте фильтры до поиска.")

    # 4. Статистика видео и каналов.
    video_ids = list(hits)
    print(f"\nЗагрузка данных о {len(video_ids)} видео...")
    videos = yt.videos(video_ids)
    channel_ids = list(dict.fromkeys(
        (v.get("snippet") or {}).get("channelId") for v in videos.values() if v
    ))
    channel_ids = [c for c in channel_ids if c]
    print(f"Загрузка данных о {len(channel_ids)} каналах...")
    channels = yt.channels(channel_ids)

    # 5. Расчёт, фильтры, сортировка.
    all_rows = build_rows(hits, queries, videos, channels, cfg)
    for r in all_rows:
        r["base_ok"] = passes_base(r, cfg)
    final_rows = sorted((r for r in all_rows if r["base_ok"] and passes_subs(r, cfg)), key=sort_key, reverse=True)
    final_ids = {r["id"] for r in final_rows}

    # 6. Проверка лучших видео на других языках.
    sources, check_results = [], []
    if cfg["translate_check"] and cfg["translate_top_videos"] > 0:
        translator = Translator(cfg["translate_provider"], cfg["api_key"], cache, yt.session)
        try:
            sources, check_results = translation_check(yt, translator, final_rows, queries, cfg, quota)
        except FinderError as exc:
            print(f"\nПроверка на других языках прервана: {exc}\nОсновной отчёт всё равно будет сохранён.")

    sheets = [
        ("Видео", make_video_table(final_rows, queries), "ratio"),
        ("Ниши", make_niche_table(all_rows, final_ids, queries, failed, cfg), None),
    ]
    if check_results:
        sheets.append(("Языки", make_languages_table(sources, check_results, cfg["translate_languages"]), "verdict"))
        sheets.append(("Переводы", make_translations_table(sources, check_results, cfg["translate_languages"]), "verdict"))
    sheets.append(("Каналы", make_channel_table(final_rows), None))
    if cfg["suggest"]:
        sheets.append(("Подсказки", make_suggest_table(suggestions), None))
    path = save_report(cfg, sheets)

    hidden = sum(1 for r in final_rows if r["hidden"])
    print(f"\nУникальных видео найдено: {len(all_rows)}, прошли фильтры: {len(final_rows)}"
          + (f" (из них со скрытыми подписчиками: {hidden})" if hidden else "") + ".")
    print(f"Каналов в отчёте: {len({r['channel_id'] for r in final_rows})}.")
    if check_results:
        free = sum(1 for r in check_results if r["verdict"] == VERDICT_FREE)
        print(f"Проверка на других языках: свободно {free} из {len(check_results)} пар видео и язык "
              "(подробно на листах «Языки» и «Переводы»).")
    print(quota.summary())
    print(f"Отчёт сохранён: {path}")


def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass
    config_path = Path(sys.argv[1]) if len(sys.argv) > 1 else BASE_DIR / "config.json"
    try:
        run(config_path)
    except FinderError as exc:
        print(f"\nОШИБКА: {exc}")
        return 1
    except KeyboardInterrupt:
        print("\nОстановлено пользователем. Скачанные данные сохранены в кэше.")
        return 130
    return 0


if __name__ == "__main__":
    sys.exit(main())
